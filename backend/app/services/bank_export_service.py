"""
Bank Export Service for Cambodia Enterprise HRMS.
Generates compliant payroll disbursement files for ABA Bank, ACLEDA Bank, and Standard format (CSV/Excel).
"""

import io
import csv
from typing import List, Dict, Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.models.payroll import PayrollItem, PayrollRun
from app.models.employee import Employee


class BankExportService:
    @staticmethod
    def generate_csv(items_with_emp: List[tuple[PayrollItem, Employee]], bank_format: str = "ABA") -> str:
        """Generates bank-specific CSV text."""
        output = io.StringIO()
        writer = csv.writer(output)

        if bank_format.upper() == "ABA":
            # ABA Bank iBanking Payroll Batch Format
            headers = [
                "Beneficiary Account",
                "Beneficiary Name",
                "Amount",
                "Currency",
                "Remarks",
                "Employee Code",
            ]
            writer.writerow(headers)
            for item, emp in items_with_emp:
                acc = emp.bank_account_number or "000000000"
                name = emp.bank_account_name or f"{emp.first_name_en} {emp.last_name_en}"
                curr = item.currency_contract or "USD"
                amt = float(item.net_salary_usd if curr == "USD" else item.net_salary_khr)
                remarks = f"Salary {item.payroll_run.period.name if item.payroll_run and item.payroll_run.period else ''}"
                writer.writerow([acc, name, f"{amt:.2f}", curr, remarks, emp.employee_code])

        elif bank_format.upper() == "ACLEDA":
            # ACLEDA Bank Corporate Payout Format
            headers = [
                "No",
                "Account Number",
                "Account Name",
                "Currency",
                "Amount",
                "Description",
                "Staff ID",
            ]
            writer.writerow(headers)
            for idx, (item, emp) in enumerate(items_with_emp, start=1):
                acc = emp.bank_account_number or "000000000000"
                name = emp.bank_account_name or f"{emp.first_name_en} {emp.last_name_en}"
                curr = item.currency_contract or "USD"
                amt = float(item.net_salary_usd if curr == "USD" else item.net_salary_khr)
                writer.writerow([idx, acc, name, curr, f"{amt:.2f}", "Monthly Salary", emp.employee_code])

        else:
            # Universal Enterprise Format
            headers = [
                "Employee Code",
                "Employee Name (EN)",
                "Employee Name (KH)",
                "Bank Name",
                "Bank Account Number",
                "Bank Account Name",
                "Currency",
                "Net Salary Amount",
                "Net KHR",
                "Net USD",
            ]
            writer.writerow(headers)
            for item, emp in items_with_emp:
                writer.writerow([
                    emp.employee_code,
                    f"{emp.first_name_en} {emp.last_name_en}",
                    f"{emp.last_name_kh} {emp.first_name_kh}",
                    emp.bank_name or "ABA Bank",
                    emp.bank_account_number or "",
                    emp.bank_account_name or "",
                    item.currency_contract,
                    f"{float(item.net_salary_usd if item.currency_contract == 'USD' else item.net_salary_khr):.2f}",
                    f"{int(float(item.net_salary_khr)):,}",
                    f"{float(item.net_salary_usd):.2f}",
                ])

        return output.getvalue()

    @staticmethod
    def generate_excel(items_with_emp: List[tuple[PayrollItem, Employee]], bank_format: str = "ABA") -> bytes:
        """Generates beautifully styled Excel sheet for bank submission & audit."""
        wb = Workbook()
        ws = wb.active
        ws.title = f"{bank_format.upper()}_Payroll"

        # Headers styling
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        headers = [
            "No",
            "Employee Code",
            "Beneficiary Name",
            "Bank Name",
            "Account Number",
            "Currency",
            "Net Payout Amount",
            "Equivalent KHR",
            "Equivalent USD",
            "Payment Status",
        ]
        ws.append(headers)

        for col_num, _ in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for idx, (item, emp) in enumerate(items_with_emp, start=1):
            curr = item.currency_contract or "USD"
            amt = float(item.net_salary_usd if curr == "USD" else item.net_salary_khr)
            row = [
                idx,
                emp.employee_code,
                emp.bank_account_name or f"{emp.first_name_en} {emp.last_name_en}",
                emp.bank_name or "ABA Bank",
                str(emp.bank_account_number or ""),
                curr,
                amt,
                int(float(item.net_salary_khr)),
                float(item.net_salary_usd),
                "PENDING DISBURSEMENT",
            ]
            ws.append(row)
            row_idx = ws.max_row
            for col_idx in range(1, len(row) + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                c.border = thin_border
                if col_idx in [7, 8, 9]:
                    c.alignment = Alignment(horizontal="right")
                    c.number_format = "#,##0.00" if col_idx != 8 else "#,##0"

        # Adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
