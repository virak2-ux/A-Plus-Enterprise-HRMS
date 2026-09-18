import io
import zipfile
import pytest
from openpyxl import load_workbook


def test_backup_excel_authenticated(client, auth_headers):
    """Verify full database backup export in multi-sheet Excel (.xlsx) format."""
    response = client.get("/api/v1/system/backup/excel", headers=auth_headers)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment; filename=" in response.headers["content-disposition"]
    assert response.headers["content-disposition"].endswith('.xlsx"')

    # Validate workbook contents using openpyxl
    wb = load_workbook(io.BytesIO(response.content))
    expected_sheets = {
        "Employees",
        "Departments",
        "Positions",
        "Attendance",
        "Leave Requests",
        "Payroll Ledger",
        "Audit Logs",
        "System Settings",
    }
    assert expected_sheets.issubset(set(wb.sheetnames))

    # Verify Employees sheet has columns and rows
    ws_emp = wb["Employees"]
    assert ws_emp.max_row >= 1
    header_cells = [cell.value for cell in ws_emp[1]]
    assert "Employee Code" in header_cells
    assert "Name (Khmer)" in header_cells
    assert "Name (English)" in header_cells
    assert "Base Salary" in header_cells


def test_backup_csv_zip_archive(client, auth_headers):
    """Verify full database backup export in compressed CSV .zip archive."""
    response = client.get("/api/v1/system/backup/csv", headers=auth_headers)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/zip"
    assert "attachment; filename=" in response.headers["content-disposition"]
    assert response.headers["content-disposition"].endswith('.zip"')

    # Validate zip file contents
    with zipfile.ZipFile(io.BytesIO(response.content), "r") as zf:
        file_list = zf.namelist()
        expected_files = [
            "employees.csv",
            "departments.csv",
            "positions.csv",
            "attendance_records.csv",
            "leave_requests.csv",
            "payroll_ledger.csv",
            "audit_logs.csv",
            "system_settings.csv",
        ]
        for f in expected_files:
            assert f in file_list
            content = zf.read(f)
            # Verify UTF-8 BOM is present
            assert content.startswith(b"\xef\xbb\xbf")


def test_backup_single_table_csv(client, auth_headers):
    """Verify single-table CSV export with Khmer UTF-8 BOM encoding."""
    response = client.get("/api/v1/system/backup/csv?table=employees", headers=auth_headers)
    assert response.status_code == 200
    assert "text/csv" in response.headers["content-type"]
    assert "attachment; filename=" in response.headers["content-disposition"]
    assert "employees" in response.headers["content-disposition"]

    # Verify content starts with UTF-8 BOM
    assert response.content.startswith(b"\xef\xbb\xbf")
    decoded_text = response.content.decode("utf-8-sig")
    assert "Employee Code" in decoded_text
    assert "Name (Khmer)" in decoded_text


def test_backup_unauthorized(client):
    """Verify unauthenticated access to backup endpoints is rejected."""
    resp_excel = client.get("/api/v1/system/backup/excel")
    assert resp_excel.status_code == 401

    resp_csv = client.get("/api/v1/system/backup/csv")
    assert resp_csv.status_code == 401
